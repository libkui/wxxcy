from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

_DEFAULT_XLSX_PATH = Path(__file__).with_name("devices.xlsx")
_DEFAULT_SHEET_NAME = "devices"

_REQUIRED_COLUMNS = ("SWITCH_IP", "USERNAME", "PASSWORD")
_ALL_COLUMNS = ("SWITCH_IP", "Device_name", "device_type", "USERNAME", "PASSWORD", "PORT")


def _normalize_header(value: Any) -> str:
    return str(value or "").strip()


def _to_int(value: Any, default: int) -> int:
    if value is None or value == "":
        return default
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default


def load_devices_from_xlsx(
    path: Path = _DEFAULT_XLSX_PATH,
    sheet_name: Optional[str] = _DEFAULT_SHEET_NAME,
) -> List[Dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(
            f"Device inventory file not found: {path}. "
            f"Please create an Excel file with columns: {', '.join(_ALL_COLUMNS)}"
        )

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active

        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [_normalize_header(cell) for cell in header_row]

        col_index: Dict[str, int] = {}
        for idx, name in enumerate(headers):
            if not name:
                continue
            col_index[name.upper()] = idx

        missing = [c for c in _REQUIRED_COLUMNS if c.upper() not in col_index]
        if missing:
            raise ValueError(
                f"Missing required columns in {path}: {', '.join(missing)}. "
                f"Found: {', '.join(h for h in headers if h)}"
            )

        devices: List[Dict[str, Any]] = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or v == "" for v in row):
                continue

            switch_ip = row[col_index["SWITCH_IP"]]
            if switch_ip is None or str(switch_ip).strip() == "":
                continue

            device: Dict[str, Any] = {
                "SWITCH_IP": str(switch_ip).strip(),
                "USERNAME": str(row[col_index["USERNAME"]] or "").strip(),
                "PASSWORD": str(row[col_index["PASSWORD"]] or "").strip(),
            }

            if "DEVICE_NAME" in col_index:
                device["Device_name"] = str(row[col_index["DEVICE_NAME"]] or "").strip()
            if "DEVICE_TYPE" in col_index:
                device["device_type"] = str(row[col_index["DEVICE_TYPE"]] or "").strip()
            if "PORT" in col_index:
                device["PORT"] = _to_int(row[col_index["PORT"]], default=22)

            if not device.get("device_type"):
                device["device_type"] = "huawei"
            if not device.get("PORT"):
                device["PORT"] = 22

            devices.append(device)

        return devices
    finally:
        workbook.close()


DEVICES = load_devices_from_xlsx()
